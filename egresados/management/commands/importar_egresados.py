import datetime
import os
import openpyxl
from django.core.management.base import BaseCommand
from django.conf import settings
from egresados.models import Carrera, Egresado, CohorteIngresantes


class Command(BaseCommand):
    help = "Importa la base de datos de egresados y cohortes desde el archivo Excel 'Banco de Datos de Egresados.xlsx'"

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default='Banco de Datos de Egresados.xlsx',
            help='Ruta al archivo Excel de origen (por defecto: Banco de Datos de Egresados.xlsx)'
        )

    def handle(self, *args, **options):
        excel_path = options['file']
        if not os.path.isabs(excel_path):
            excel_path = os.path.join(settings.BASE_DIR, excel_path)

        if not os.path.exists(excel_path):
            self.stderr.write(self.style.ERROR(f"No se encontró el archivo: {excel_path}"))
            return

        self.stdout.write(self.style.NOTICE(f"Cargando archivo: {excel_path}..."))
        wb = openpyxl.load_workbook(excel_path, data_only=True)

        # 1. Catálogo inicial de carreras
        carreras_default = [
            'Ingeniería en Electricidad',
            'Ingeniería en Electrónica',
            'Ingeniería en Sistemas Informáticos',
            'Ingeniería Civil',
        ]
        carreras_map = {}
        for nombre_carrera in carreras_default:
            carrera_obj, _ = Carrera.objects.get_or_create(
                nombre=nombre_carrera,
                defaults={'duracion_teorica': 5, 'activa': True}
            )
            carreras_map[nombre_carrera] = carrera_obj

        # 2. Ingesta de Hoja 1 (FCyT) - Egresados
        if 'FCyT' in wb.sheetnames:
            ws_egresados = wb['FCyT']
            creados = 0
            omitidos = 0
            actualizados = 0

            # Los registros empiezan en la fila 4 (fila 1: vacío, fila 2: encabezados, fila 3: filtros)
            for row in range(4, ws_egresados.max_row + 1):
                num_val = ws_egresados.cell(row, 1).value
                if num_val is None or str(num_val).strip() == '':
                    continue

                try:
                    num_egresado = int(float(num_val))
                except (ValueError, TypeError):
                    continue

                nombres = str(ws_egresados.cell(row, 2).value or '').strip()
                apellidos = str(ws_egresados.cell(row, 3).value or '').strip()
                cin_raw = ws_egresados.cell(row, 4).value
                cin = str(int(float(cin_raw))) if isinstance(cin_raw, (int, float)) else str(cin_raw or '').strip()

                fecha_nac_val = ws_egresados.cell(row, 5).value
                fecha_nac = None
                if isinstance(fecha_nac_val, (datetime.datetime, datetime.date)):
                    fecha_nac = fecha_nac_val if isinstance(fecha_nac_val, datetime.date) else fecha_nac_val.date()

                lugar_nac = str(ws_egresados.cell(row, 6).value or '').strip()
                sexo_raw = str(ws_egresados.cell(row, 7).value or '').strip().upper()
                sexo = 'M' if sexo_raw.startswith('M') else 'F'

                nombre_carrera = str(ws_egresados.cell(row, 8).value or '').strip()
                carrera_obj = carreras_map.get(nombre_carrera)
                if not carrera_obj:
                    carrera_obj, _ = Carrera.objects.get_or_create(
                        nombre=nombre_carrera,
                        defaults={'duracion_teorica': 5, 'activa': True}
                    )
                    carreras_map[nombre_carrera] = carrera_obj

                ingreso_val = ws_egresados.cell(row, 9).value
                anio_ingreso = int(float(ingreso_val)) if ingreso_val else 0

                egreso_val = ws_egresados.cell(row, 10).value
                anio_egreso = int(float(egreso_val)) if egreso_val else 0

                link_repo = ws_egresados.cell(row, 11).value
                link_repo = str(link_repo).strip() if link_repo else None

                # Verificamos si ya existe por número de egresado (clave única / PK)
                if Egresado.objects.filter(numero_egresado=num_egresado).exists():
                    omitidos += 1
                else:
                    Egresado.objects.create(
                        numero_egresado=num_egresado,
                        nombres=nombres,
                        apellidos=apellidos,
                        cin=cin,
                        fecha_nacimiento=fecha_nac,
                        lugar_nacimiento=lugar_nac,
                        sexo=sexo,
                        carrera=carrera_obj,
                        anio_ingreso=anio_ingreso,
                        anio_egreso=anio_egreso,
                        link_repositorio=link_repo,
                    )
                    creados += 1

            self.stdout.write(
                self.style.SUCCESS(
                    f"[FCyT - Egresados] Procesamiento completado: {creados} creados, {omitidos} omitidos por duplicado."
                )
            )
        else:
            self.stderr.write(self.style.WARNING("Hoja 'FCyT' no encontrada."))

        # 3. Ingesta de Hoja 'totales' - Cohortes de Ingresantes
        if 'totales' in wb.sheetnames:
            ws_totales = wb['totales']
            cohortes_creadas = 0

            # Mapeo de columnas según fila 36 y 37
            # Col 3: Electricidad, Col 6: Electrónica, Col 9: Informática, Col 12: Civil
            carreras_columnas = [
                (3, 'Ingeniería en Electricidad'),
                (6, 'Ingeniería en Electrónica'),
                (9, 'Ingeniería en Sistemas Informáticos'),
                (12, 'Ingeniería Civil'),
            ]

            for row in range(38, 52):
                anio_val = ws_totales.cell(row, 2).value
                if not anio_val:
                    continue
                try:
                    anio_cohorte = int(float(anio_val))
                except (ValueError, TypeError):
                    continue

                for col_idx, nombre_carrera in carreras_columnas:
                    carrera_obj = carreras_map.get(nombre_carrera)
                    if not carrera_obj:
                        continue

                    ingresantes_val = ws_totales.cell(row, col_idx).value
                    if ingresantes_val is not None:
                        try:
                            cantidad = int(float(ingresantes_val))
                            CohorteIngresantes.objects.update_or_create(
                                carrera=carrera_obj,
                                anio_cohorte=anio_cohorte,
                                defaults={'cantidad_ingresantes': cantidad}
                            )
                            cohortes_creadas += 1
                        except (ValueError, TypeError):
                            pass

            self.stdout.write(
                self.style.SUCCESS(
                    f"[Totales - Cohortes] Procesadas {cohortes_creadas} registros de ingresantes por cohorte."
                )
            )
        else:
            self.stderr.write(self.style.WARNING("Hoja 'totales' no encontrada."))
